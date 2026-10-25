import streamlit as st
import pandas as pd
import numpy as np
import plotly.express as px
import io

st.set_page_config(page_title="Marketing 9-Box Chart", layout="wide")

# ==========================================
# HEADER
# ==========================================
st.markdown("""
    <div style="background: linear-gradient(135deg, #1e293b 0%, #0f172a 100%); 
                padding: 1.2rem; border-radius: 10px; margin-bottom: 1rem; color: white;
                text-align: center;">
        <h2 style="margin:0; color: #38bdf8; font-weight: 800;">
            📊 Marketing Performance 9-Box Chart
        </h2>
        <p style="margin: 0.2rem 0 0 0; opacity: 0.9; font-size: 1rem;">
            Strategic Financial Output vs People Score (Proportional Allocation Matrix)
        </p>
    </div>
""", unsafe_allow_html=True)

# ==========================================
# FILE UPLOADER & DATA LOADING
# ==========================================
st.markdown("### 📂 Upload Source Data")
uploaded_file = st.file_uploader("Upload Excel file (e.g., Financial_Metric_Marketing_Dummy.xlsx)", type=["xlsx"])

if uploaded_file is None:
    st.info("☝️ Please upload your Excel data file to begin.")
    st.stop()

try:
    xls = pd.ExcelFile(uploaded_file)

    req_sheets = ['PL_Local', 'Matrix']
    missing_sheets = [s for s in req_sheets if s not in xls.sheet_names]
    if missing_sheets:
        st.error(f"⚠️ The following sheets were not found in the Excel file: {', '.join(missing_sheets)}")
        st.stop()

    # 1. READ PL_LOCAL (Raw Data)
    df_pl = pd.read_excel(xls, sheet_name='PL_Local')


    # Enhanced Currency Cleaner to handle accounting formats
    def clean_curr(x):
        if pd.isna(x): return 0.0
        if isinstance(x, (int, float)): return float(x)
        s = str(x).strip()
        if s.startswith('(') and s.endswith(')'):
            s = '-' + s[1:-1]
        s = s.replace(' ', '').replace('.', '').replace(',', '.')
        try:
            return float(s)
        except:
            return 0.0


    col_gs = 'Gross Sales' if 'Gross Sales' in df_pl.columns else None
    col_ret = 'Sales Return' if 'Sales Return' in df_pl.columns else ('Return' if 'Return' in df_pl.columns else None)
    col_cm = 'Contribution Margin'
    col_op = 'Operating Income' if 'Operating Income' in df_pl.columns else 'Operating Profit'

    # Calculate actual baseline figures for the entire population (Including zero and negative/return values)
    if col_gs and col_ret:
        df_pl['Calc_Sales'] = df_pl[col_gs].apply(clean_curr) - df_pl[col_ret].apply(clean_curr).abs()
    elif col_gs:
        df_pl['Calc_Sales'] = df_pl[col_gs].apply(clean_curr)
    else:
        df_pl['Calc_Sales'] = 0.0

    df_pl[col_cm] = df_pl[col_cm].apply(clean_curr) if col_cm in df_pl.columns else 0.0
    df_pl[col_op] = df_pl[col_op].apply(clean_curr) if col_op in df_pl.columns else 0.0

    # Store standard lowercase strings for safe comparison
    df_pl['Match_Brand'] = df_pl['Brand'].astype(str).str.strip().str.lower()
    df_pl['Match_Category'] = df_pl['Category'].astype(str).str.strip().str.lower()
    df_pl['Match_Pillar'] = df_pl['Pillar'].astype(str).str.strip().str.lower()

    # Calculate baseline totals for the control check
    total_pop_sales = df_pl['Calc_Sales'].sum()
    total_pop_cm = df_pl[col_cm].sum()
    total_pop_op = df_pl[col_op].sum()

    # 2. FIND PEOPLE SCORES (Normal & Adjusted)
    score_map_normal = {}
    score_map_adj = {}

    if 'Head-Count' in xls.sheet_names:
        df_hc = pd.read_excel(xls, sheet_name='Head-Count')

        col_normal, col_adj = None, None
        for c in df_hc.columns:
            cl = str(c).lower()
            if 'people score' in cl and 'adjusted' not in cl:
                col_normal = c
            if 'adjusted' in cl and 'people score' in cl:
                col_adj = c

        if col_normal:
            score_map_normal = dict(
                zip(df_hc['Name'].astype(str).str.strip(), pd.to_numeric(df_hc[col_normal], errors='coerce')))
        if col_adj:
            score_map_adj = dict(
                zip(df_hc['Name'].astype(str).str.strip(), pd.to_numeric(df_hc[col_adj], errors='coerce')))

    # ==========================================
    # DASHBOARD CONFIGURATION UI
    # ==========================================
    st.markdown("---")
    st.markdown("### ⚙️ Dashboard Configuration")

    col_cfg1, col_cfg2, col_cfg3 = st.columns(3)
    with col_cfg1:
        # X-Axis is now Financial Output
        x_metric = st.selectbox(
            "X-Axis Metric (Financial Output):",
            ["Contribution Margin", "Operating Profit", "Adjusted Contribution Margin"]
        )
    with col_cfg2:
        # Y-Axis is now People Score
        y_metric = st.selectbox(
            "Y-Axis Metric (People Score):",
            ["People Score (1-3)", "Adjusted People Score (1-10)"],
            help="Choose the granularity of the workload and performance evaluation."
        )
    with col_cfg3:
        shared_service_pct = st.number_input(
            "Shared Service Overhead Allocation (%)",
            min_value=0.0, max_value=100.0, value=10.0, step=1.0,
            help="Percentage of the Total Active Marketing Scope allocated to the 'Shared Service' team."
        ) / 100.0

    # COLLAPSIBLE README METHODOLOGY
    with st.expander("📖 View Engine Methodology & Rules (README)"):
        st.markdown("""
        ### 📊 Marketing Performance 9-Box Data Engine

        To ensure high data integrity, prevent financial double-counting, and reflect accurate business impact, the data engine follows strict methodological rules when parsing the `PL_Local` and `Matrix` mapping sheets:

        **1. Strict Master Brand Anchor (Zero Cross-Pillar Leakage)**
        A common data issue arises when drawing data solely based on broad Categories or Pillars because they may contain items managed by other divisions (e.g., the "Others" pillar).
        *Rule:* The engine utilizes a **Master Brand Filter**. It extracts only the explicit Brands listed in the `Matrix`. Any row in the `PL_Local` raw dataset belonging to a brand *not* explicitly listed is immediately **scoped out** and excluded. 

        **2. Pillar & Category Inclusion via Master Brand Intersection**
        Personnel mapped strictly by "Pillar" or "Category" (without explicit brands) act as "Passengers" in the data distribution.
        *Rule:* They will only receive a financial allocation IF the specific Brand inside their Pillar/Category successfully passes the Master Brand whitelist check (meaning the brand is explicitly managed by someone else on the team). This ensures they receive credit for relevant products without accidentally absorbing unmanaged items from the wider P&L population.

        **3. Absolute Population Inclusion (Returns act as Deductions)**
        *Rule:* All mapped items are evaluated exactly as they appear in the raw P&L data. Negative values and Sales Returns are mathematically forced to act as absolute deductions (`Gross - absolute(Return)`) ensuring the final output perfectly mirrors the actual P&L net financial reality without artificially inflating sales.

        **4. Shared Service Overhead Allocation**
        Administrative and shared service personnel do not handle specific brands but support the entire marketing department.
        *Rule:* The engine calculates the Grand Total Active Marketing Scope. A configurable percentage (default 10%) of this grand total forms the "Shared Service Overhead". This pool is then divided equally among all personnel explicitly designated with the "Shared Service" pillar label in the `Matrix` sheet.

        **5. Proportional Headcount Allocation (Anti-Double Counting)**
        When multiple core marketing personnel are assigned to the same Brand, Category, or Pillar, standard summation would lead to inflated, double-counted financial metrics.
        *Rule:* After deducting the Shared Service Overhead (Step 4), the remaining value of any given transaction row is equally divided among the `N` core marketing personnel actively handling it. Rows not mapped to any personnel are excluded entirely, ensuring a 100% accurate financial reflection without double counting.

        **6. Dynamic Y-Axis Scaling (Adjusted People Score 1-10)**
        Users can toggle between the standard `People Score (1-3)` and the `Adjusted People Score (1-10)`. 
        *Rule:* The adjusted score is calculated using a continuous mathematical formula integrating PA averages and WLA multipliers: `=ROUND((PA Average/5)*7 + WLA_Bonus, 2)`. Where WLA Low = 0, Mid = 1.5, High = 3.0. This prevents visual overlapping on the scatter plot and provides granular performance differentiation.

        **7. Operating Profit Cap (Adjusted Contribution Margin)**
        Contribution Margin is utilized as the primary financial indicator for X-Axis placement. However, evaluating solely on gross margin can overlook severe operational inefficiencies.
        *Rule:* If the "Adjusted Contribution Margin" metric is selected, the engine scans the net Operating Profit. If a personnel's Operating Profit is strictly negative (< 0), their Adjusted Contribution Margin is mathematically hard-capped just below the Low-Med X-Axis threshold. This ensures personnel generating operational losses are automatically reclassified into the bottom financial tier regardless of how large their gross margin is.
        """)

    # 3. READ MATRIX & BUILD STRICT MARKETING SCOPE
    df_mat_raw = pd.read_excel(xls, sheet_name='Matrix', header=None)
    persons = []
    shared_service_persons = []
    global_active_brands = set()

    for idx, row in df_mat_raw.iterrows():
        try:
            no = int(row[0])
        except:
            continue

        name = str(row[1]).strip()
        if pd.isna(name) or name.lower() in ['nan', '']: continue

        # Extract Scores
        ps_normal = score_map_normal.get(name, 0.0)
        ps_adj = score_map_adj.get(name, 0.0)

        # Fallbacks
        if pd.isna(ps_normal) or ps_normal == 0.0:
            if len(row) > 15 and pd.notna(row[15]):
                try:
                    ps_normal = float(row[15])
                except:
                    ps_normal = 0.0

        if pd.isna(ps_adj) or ps_adj == 0.0:
            ps_adj = ps_normal  # Fallback to standard if adjusted is entirely missing

        # Extract Mapping Constraints
        pillar = str(row[2]).strip().lower() if pd.notna(row[2]) else ""
        brands = [str(x).strip().lower() for x in row[3:9] if pd.notna(x) and str(x).strip() != '']
        categories = [str(x).strip().lower() for x in row[9:15] if pd.notna(x) and str(x).strip() != '']

        if pillar == 'shared service':
            shared_service_persons.append({'name': name, 'score_normal': ps_normal, 'score_adj': ps_adj})
        else:
            clean_brands = list(set([b for b in brands if b != 'nan']))
            clean_cats = list(set([c for c in categories if c != 'nan']))

            # Anchor ALL mapped brands into the Master Brand list
            global_active_brands.update(clean_brands)

            persons.append({
                'name': name,
                'score_normal': ps_normal,
                'score_adj': ps_adj,
                'pillar': pillar,
                'brands': clean_brands,
                'categories': clean_cats
            })

    # Calculate Explicit Master Brand Scope (For Validation Use Only)
    mask_explicit = df_pl['Match_Brand'].isin(global_active_brands)
    df_explicit = df_pl[mask_explicit]
    explicit_sales = df_explicit['Calc_Sales'].sum()

    # 4. PROPORTIONAL CALCULATION (Strict Master Brand Anchor & Anti-Leakage)
    person_totals = {
        p['name']: {'Sales': 0.0, 'CM': 0.0, 'OP': 0.0, 'Score_Normal': p['score_normal'], 'Score_Adj': p['score_adj']}
        for p in persons}

    total_marketing_sales = 0.0
    total_marketing_cm = 0.0
    total_marketing_op = 0.0
    breakdown_data = []

    # Core multiplier ensures the total sum does not inflate beyond actuals
    core_multiplier = 1.0
    if shared_service_persons:
        core_multiplier = 1.0 - shared_service_pct

    for idx, row in df_pl.iterrows():
        sales = row['Calc_Sales']
        row_brand = row['Match_Brand']
        row_category = row['Match_Category']
        row_pillar = row['Match_Pillar']

        # SCOPE-OUT: Master Brand Filter. Exclude brands not explicitly handled by marketing.
        if row_brand not in global_active_brands:
            continue

        # Determine which marketing personnel handle this specific row
        covered_by = []
        for p in persons:
            is_match = False
            if p['pillar'] and p['pillar'] == row_pillar:
                is_match = True
            elif row_brand in p['brands']:
                is_match = True
            elif row_category in p['categories']:
                is_match = True

            if is_match:
                covered_by.append(p['name'])

        N = len(covered_by)

        # SCOPE-IN: Distribute proportionally among the assigned personnel
        if N > 0:
            cm = row[col_cm]
            op = row[col_op]

            total_marketing_sales += sales
            total_marketing_cm += cm
            total_marketing_op += op

            for name in covered_by:
                person_totals[name]['Sales'] += (sales * core_multiplier) / N
                person_totals[name]['CM'] += (cm * core_multiplier) / N
                person_totals[name]['OP'] += (op * core_multiplier) / N

            # Audit Trail for Active Scope Breakdown (SALES ONLY)
            ss_allocation = sales * shared_service_pct if shared_service_persons else 0.0
            core_allocation = sales * core_multiplier

            breakdown_data.append({
                'Brand': row['Brand'],
                'Category': row['Category'],
                'PL Original Sales (Gross - Return)': sales,
                'System Core Allocation': core_allocation,
                'Shared Service Allocation': ss_allocation,
                'Covered By': ", ".join(sorted(covered_by))
            })

    # 5. ALLOCATE SHARED SERVICE POOL
    if shared_service_persons:
        ss_count = len(shared_service_persons)
        ss_sales = (total_marketing_sales * shared_service_pct) / ss_count
        ss_cm = (total_marketing_cm * shared_service_pct) / ss_count
        ss_op = (total_marketing_op * shared_service_pct) / ss_count

        for p in shared_service_persons:
            person_totals[p['name']] = {
                'Sales': ss_sales,
                'CM': ss_cm,
                'OP': ss_op,
                'Score_Normal': p['score_normal'],
                'Score_Adj': p['score_adj']
            }

    # 6. DATA INTEGRITY & BREAKDOWN CHECKS
    st.markdown("---")
    st.markdown("### ⚖️ Data Integrity & Control Check")
    st.caption(
        "Validates data allocation by comparing the distributed Marketing Team totals against the Total PL Population. Unmanaged brands are safely scoped out.")

    pct_sales = (total_marketing_sales / total_pop_sales * 100) if total_pop_sales else 0
    pct_cm = (total_marketing_cm / total_pop_cm * 100) if total_pop_cm else 0
    pct_op = (total_marketing_op / total_pop_op * 100) if total_pop_op else 0

    check_df = pd.DataFrame({
        'Financial Metric': ['Sales (Gross - Return)', 'Contribution Margin', 'Operating Profit'],
        'Total PL Population': [total_pop_sales, total_pop_cm, total_pop_op],
        'Marketing Team Dist. Total': [total_marketing_sales, total_marketing_cm, total_marketing_op],
        '% Marketing Contribution': [pct_sales, pct_cm, pct_op],
        'Scoped Out (Others/No Brand)': [total_pop_sales - total_marketing_sales, total_pop_cm - total_marketing_cm,
                                         total_pop_op - total_marketing_op]
    })

    st.dataframe(
        check_df.style.format({
            'Total PL Population': 'Rp {:,.0f}',
            'Marketing Team Dist. Total': 'Rp {:,.0f}',
            '% Marketing Contribution': '{:.2f}%',
            'Scoped Out (Others/No Brand)': 'Rp {:,.0f}'
        }),
        use_container_width=True, hide_index=True
    )

    # Internal validation alert (invisible if correct)
    if not np.isclose(explicit_sales, total_marketing_sales, atol=1):
        st.warning(
            "⚠️ Control Check Notice: Discrepancy detected between explicit master brand filter and final distribution.")

    # ---------------------------------------------------------
    # ACTIVE SCOPE BREAKDOWN (SALES AUDIT TRAIL) WITH EXPANDER
    # ---------------------------------------------------------
    with st.expander("🔍 View Active Scope Breakdown (Sales Traceability)"):
        st.caption(
            "This table isolates the Sales figures per Brand and Category to pinpoint variances. It maps the original P&L Sales against the system allocations (Core vs. Shared Service) to ensure zero Rupiah is lost in the calculation.")

        if breakdown_data:
            df_bd = pd.DataFrame(breakdown_data)

            # Aggregate identical rows
            df_bd_agg = df_bd.groupby(['Brand', 'Category', 'Covered By'], as_index=False).sum()

            # Calculate Variance (Should exactly equal 0)
            df_bd_agg['Variance (Unallocated)'] = df_bd_agg['PL Original Sales (Gross - Return)'] - (
                        df_bd_agg['System Core Allocation'] + df_bd_agg['Shared Service Allocation'])

            df_bd_agg = df_bd_agg.sort_values(by='PL Original Sales (Gross - Return)', ascending=False)

            st.dataframe(
                df_bd_agg.style.format({
                    'PL Original Sales (Gross - Return)': 'Rp {:,.0f}',
                    'System Core Allocation': 'Rp {:,.0f}',
                    'Shared Service Allocation': 'Rp {:,.0f}',
                    'Variance (Unallocated)': 'Rp {:,.0f}'
                }),
                use_container_width=True, hide_index=True
            )

            # Download button for Breakdown
            buffer_bd = io.BytesIO()
            with pd.ExcelWriter(buffer_bd, engine='openpyxl') as writer:
                df_bd_agg.to_excel(writer, index=False, sheet_name='Sales_Breakdown')

            st.download_button(
                label="📥 Download Sales Breakdown (Excel)",
                data=buffer_bd.getvalue(),
                file_name="Marketing_Sales_Breakdown.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="secondary",
                key="btn_dl_breakdown"
            )
        else:
            st.info("No active brand or category data found.")

    # 7. COMPILE FINAL RESULTS (Apply chosen Y-Metric dynamically)
    results = []
    for name, data in person_totals.items():
        # Inject the correct score depending on user selection for the Y-Axis
        chosen_score = data['Score_Adj'] if "Adjusted" in y_metric else data['Score_Normal']

        results.append({
            'Name': name,
            y_metric: chosen_score,
            'Sales': data['Sales'],
            'Contribution Margin': data['CM'],
            'Operating Profit': data['OP']
        })

    df_res = pd.DataFrame(results)

    if df_res.empty:
        st.warning("⚠️ No valid marketing personnel were processed. Please check your Matrix mapping.")
        st.stop()

    total_persons = len(df_res)

    # Initialize Adjusted Contribution Margin (Base is CM)
    df_res['Adjusted Contribution Margin'] = df_res['Contribution Margin']

    # ==========================================
    # SETTINGS: 9-BOX MIDPOINT THRESHOLDS
    # ==========================================
    st.markdown("---")
    st.markdown("### 🎚️ 9-Box Midpoint Thresholds")
    st.info(
        f"💡 **X-Axis** is mapping to **{x_metric}**. **Y-Axis** is mapping to **{y_metric}**. Bubble size visualizes Sales.")

    # Dynamic default thresholds based on scale for Y-Axis (People Score)
    if "Adjusted" in y_metric:
        y_low_default = 4.0
        y_high_default = 7.0
    else:
        y_low_default = 1.5
        y_high_default = 2.5

    # X-Axis is now Financial Output
    mid_x = df_res[x_metric].mean() if df_res[x_metric].sum() > 0 else 1000000.0
    step_x = 1000000.0

    col_t1, col_t2, col_t3, col_t4 = st.columns(4)
    with col_t1:
        x_low = st.number_input(f"X-Axis Low-Med ({x_metric})", value=float(mid_x * 0.8), step=step_x)
    with col_t2:
        x_high = st.number_input(f"X-Axis Med-High ({x_metric})", value=float(mid_x * 1.2), step=step_x)
    with col_t3:
        y_low = st.number_input(f"Y-Axis Low-Med (People Score)", value=float(y_low_default), step=0.5)
    with col_t4:
        y_high = st.number_input(f"Y-Axis Med-High (People Score)", value=float(y_high_default), step=0.5)

    # ---------------------------------------------------------
    # APPLY OPERATING PROFIT HARD CAP (PENALTY) ON X-AXIS
    # ---------------------------------------------------------
    if x_metric == "Adjusted Contribution Margin":
        # Force any personnel with negative OP to sit directly below the x_low threshold line
        penalty_cap_value = x_low - 1
        df_res.loc[df_res['Operating Profit'] < 0, x_metric] = df_res.loc[
            df_res['Operating Profit'] < 0, x_metric].clip(upper=penalty_cap_value)

    # ==========================================
    # ASSIGN 9-BOX CATEGORY (STRICT ALIGNMENT)
    # ==========================================
    cat_protect = "B9 - Protect & scale"
    cat_keep = "B8 - Keep & strengthen"
    cat_move = "B7 - Move to higher output"
    cat_share = "B6 - Share across brands"
    cat_simplify = "B5 - Simplify the work"
    cat_reduce_some = "B4 - Reduce some"
    cat_upskill = "B3 - Upskill or replace"
    cat_replace = "B2 - Replace with tools"
    cat_reduce_most = "B1 - Reduce most (over half)"


    def assign_9box(df_target):
        # Y is People Score, X is Financial Output
        conds = [
            (df_target[y_metric] > y_high) & (df_target[x_metric] > x_high),  # High Y, High X -> B9
            (df_target[y_metric] > y_high) & (df_target[x_metric] >= x_low) & (df_target[x_metric] <= x_high),
            # High Y, Mid X -> B8
            (df_target[y_metric] > y_high) & (df_target[x_metric] < x_low),  # High Y, Low X -> B7

            (df_target[y_metric] >= y_low) & (df_target[y_metric] <= y_high) & (df_target[x_metric] > x_high),
            # Mid Y, High X -> B6
            (df_target[y_metric] >= y_low) & (df_target[y_metric] <= y_high) & (df_target[x_metric] >= x_low) & (
                        df_target[x_metric] <= x_high),  # Mid Y, Mid X -> B5
            (df_target[y_metric] >= y_low) & (df_target[y_metric] <= y_high) & (df_target[x_metric] < x_low),
            # Mid Y, Low X -> B4

            (df_target[y_metric] < y_low) & (df_target[x_metric] > x_high),  # Low Y, High X -> B3
            (df_target[y_metric] < y_low) & (df_target[x_metric] >= x_low) & (df_target[x_metric] <= x_high),
            # Low Y, Mid X -> B2
            (df_target[y_metric] < y_low) & (df_target[x_metric] < x_low)  # Low Y, Low X -> B1
        ]
        choices = [
            cat_protect,  # B9
            cat_keep,  # B8
            cat_move,  # B7
            cat_share,  # B6
            cat_simplify,  # B5
            cat_reduce_some,  # B4
            cat_upskill,  # B3
            cat_replace,  # B2
            cat_reduce_most  # B1
        ]
        return np.select(conds, choices, default=cat_simplify)


    df_res['9-Box Category'] = assign_9box(df_res)
    df_res['Bubble_Size'] = df_res['Sales'].abs().replace(0, 1)

    # ==========================================
    # SCATTER PLOT VISUALIZATION
    # ==========================================
    color_map = {
        cat_protect: '#a7c957',  # Green
        cat_keep: '#c5e1a5',  # Light Green
        cat_move: '#a8dadc',  # Light Blue
        cat_share: '#ffd166',  # Yellow/Orange
        cat_simplify: '#ffe66d',  # Yellow
        cat_reduce_some: '#f4a261',  # Orange
        cat_upskill: '#f4a261',  # Orange
        cat_replace: '#e76f51',  # Dark Orange
        cat_reduce_most: '#d62828'  # Dark Red
    }

    # Dynamic Hover Data
    hover_dict = {'Bubble_Size': False, 'Sales': ':,.0f'}
    if x_metric == "Adjusted Contribution Margin":
        hover_dict['Contribution Margin'] = ':,.0f'
        hover_dict['Operating Profit'] = ':,.0f'
        hover_dict['Adjusted Contribution Margin'] = ':,.0f'
    else:
        hover_dict['Contribution Margin'] = ':,.0f'
        hover_dict['Operating Profit'] = ':,.0f'

    fig = px.scatter(
        df_res,
        x=x_metric,
        y=y_metric,
        color='9-Box Category',
        color_discrete_map=color_map,
        size='Bubble_Size',
        hover_name='Name',
        hover_data=hover_dict,
        title=f"9-Box Matrix: {y_metric} vs {x_metric} (Bubble Size: Sales)",
        size_max=40,
        render_mode='svg'
    )

    fig.add_vline(x=x_low, line_dash="dash", line_color="black", opacity=0.7)
    fig.add_vline(x=x_high, line_dash="dash", line_color="black", opacity=0.7)
    fig.add_hline(y=y_low, line_dash="dash", line_color="black", opacity=0.7)
    fig.add_hline(y=y_high, line_dash="dash", line_color="black", opacity=0.7)

    # ---------------------------------------------------------
    # ANNOTATION (LABEL & COUNT) IN THE CENTER OF EACH QUADRANT
    # ---------------------------------------------------------
    x_min, x_max = df_res[x_metric].min(), df_res[x_metric].max()
    y_min, y_max = df_res[y_metric].min(), df_res[y_metric].max()

    x_min_plot = min(x_min, x_low) - (abs(x_max - x_min) * 0.1 if x_max != x_min else 1)
    x_max_plot = max(x_max, x_high) + (abs(x_max - x_min) * 0.1 if x_max != x_min else 1)
    y_min_plot = min(y_min, y_low) - (abs(y_max - y_min) * 0.1 if y_max != y_min else 1000)
    y_max_plot = max(y_max, y_high) + (abs(y_max - y_min) * 0.1 if y_max != y_min else 1000)

    mid_x_low = (x_min_plot + x_low) / 2
    mid_x_mid = (x_low + x_high) / 2
    mid_x_high = (x_high + x_max_plot) / 2

    mid_y_low = (y_min_plot + y_low) / 2
    mid_y_mid = (y_low + y_high) / 2
    mid_y_high = (y_high + y_max_plot) / 2

    box_counts = df_res['9-Box Category'].value_counts().to_dict()

    # Annotation coordinates (x maps to Financial, y maps to People Score)
    annotations_data = [
        {'cat': cat_move, 'x': mid_x_low, 'y': mid_y_high},  # B7
        {'cat': cat_keep, 'x': mid_x_mid, 'y': mid_y_high},  # B8
        {'cat': cat_protect, 'x': mid_x_high, 'y': mid_y_high},  # B9
        {'cat': cat_reduce_some, 'x': mid_x_low, 'y': mid_y_mid},  # B4
        {'cat': cat_simplify, 'x': mid_x_mid, 'y': mid_y_mid},  # B5
        {'cat': cat_share, 'x': mid_x_high, 'y': mid_y_mid},  # B6
        {'cat': cat_reduce_most, 'x': mid_x_low, 'y': mid_y_low},  # B1
        {'cat': cat_replace, 'x': mid_x_mid, 'y': mid_y_low},  # B2
        {'cat': cat_upskill, 'x': mid_x_high, 'y': mid_y_low}  # B3
    ]

    for item in annotations_data:
        count = box_counts.get(item['cat'], 0)
        person_text = "Person" if count == 1 else "Persons"
        # Label is always rendered, even if 0 persons, to prevent UI layout jumps
        fig.add_annotation(
            x=item['x'], y=item['y'],
            xref="x", yref="y",
            text=f"<span style='color:#1f2937;'><b>{item['cat']}</b></span><br><b>{count} {person_text}</b>",
            showarrow=False,
            font=dict(size=12, color="#374151"),
            bgcolor="rgba(255, 255, 255, 0.85)",
            bordercolor="rgba(15, 23, 42, 0.3)",
            borderwidth=1, borderpad=4
        )

    # ADD TOTAL OVERALL ANNOTATION
    total_person_text = "Person" if total_persons == 1 else "Persons"
    fig.add_annotation(
        x=0.98, y=0.98, xref="paper", yref="paper",
        text=f"<b>TOTAL:<br>{total_persons} {total_person_text}</b>",
        showarrow=False,
        font=dict(size=13, color="white"),
        bgcolor="#374151",
        bordercolor="black",
        borderwidth=1, borderpad=5
    )

    fig.update_xaxes(range=[x_min_plot, x_max_plot])
    fig.update_yaxes(range=[y_min_plot, y_max_plot])

    fig.update_layout(
        height=650,
        hovermode="closest",
        margin=dict(t=50, b=50, l=50, r=50),
        legend=dict(orientation="h", yanchor="top", y=-0.15, xanchor="center", x=0.5, title=None)
    )

    st.plotly_chart(fig, use_container_width=True, config={'displaylogo': False,
                                                           'toImageButtonOptions': {'format': 'png',
                                                                                    'filename': '9box_chart',
                                                                                    'scale': 2}})
    st.caption("📷 *Tip: Click the camera icon in the top right corner of the chart to download the image (PNG).*")

    # ==========================================
    # HTML 9-BOX VISUAL TABLE (FIXED AXES ALIGNMENT)
    # ==========================================
    st.markdown("---")
    st.markdown("### 🗂️ Strategic 9-Box Matrix Summary")

    # Shorten metric names for horizontal HTML label to prevent clutter
    if x_metric == "Adjusted Contribution Margin":
        short_x = "Adj. CM"
    elif x_metric == "Contribution Margin":
        short_x = "CM"
    elif x_metric == "Operating Profit":
        short_x = "OP"
    else:
        short_x = x_metric

    html_matrix = f"""
    <div style="font-family: Arial, sans-serif; border: 1px solid #ddd; border-collapse: collapse; width: 100%; text-align: center; margin-bottom: 15px;">
        <div style="display: flex; background-color: #1e3a5f; color: white; padding: 10px; font-weight: bold;">
            <div style="flex: 1; display: flex; align-items: center; justify-content: center; background-color: #0f172a; border-radius: 5px;">TOTAL: {total_persons} {total_person_text.upper()}</div>
            <div style="flex: 3; display: flex; align-items: center; justify-content: center;">FINANCIAL OUTPUT ({short_x})</div>
        </div>
        <div style="display: flex; background-color: #e6e6e6; font-weight: bold; border-bottom: 2px solid white;">
            <div style="flex: 1; background-color: #1e3a5f; color: white; display: flex; align-items: center; justify-content: center; writing-mode: vertical-rl; transform: rotate(180deg); text-align: center; padding: 10px 0;">
                PEOPLE SCORE<br>({y_metric})
            </div>
            <div style="flex: 3; display: flex; flex-direction: column;">
                <div style="display: flex; background-color: #e6e6e6; padding: 5px 0;">
                    <div style="flex: 1;">LOW</div>
                    <div style="flex: 1;">MID</div>
                    <div style="flex: 1;">HIGH</div>
                </div>

                <!-- TOP ROW: HIGH PEOPLE SCORE -->
                <div style="display: flex; height: 120px;">
                    <div style="flex: 0.3; background-color: #dbe4f0; display: flex; align-items: center; justify-content: center; font-weight: bold; border-right: 2px solid white; border-bottom: 2px solid white;">HIGH</div>
                    <div style="flex: 1; background-color: {color_map[cat_move]}; padding: 10px; border-right: 2px solid white; border-bottom: 2px solid white; display: flex; flex-direction: column; justify-content: center;">
                        <span style="font-size: 18px; font-weight: bold;">{cat_move}</span>
                        <span style="font-size: 24px; font-weight: 900; margin-top: 5px;">{box_counts.get(cat_move, 0)}</span>
                    </div>
                    <div style="flex: 1; background-color: {color_map[cat_keep]}; padding: 10px; border-right: 2px solid white; border-bottom: 2px solid white; display: flex; flex-direction: column; justify-content: center;">
                        <span style="font-size: 18px; font-weight: bold;">{cat_keep}</span>
                        <span style="font-size: 24px; font-weight: 900; margin-top: 5px;">{box_counts.get(cat_keep, 0)}</span>
                    </div>
                    <div style="flex: 1; background-color: {color_map[cat_protect]}; padding: 10px; border-bottom: 2px solid white; display: flex; flex-direction: column; justify-content: center;">
                        <span style="font-size: 18px; font-weight: bold;">{cat_protect}</span>
                        <span style="font-size: 24px; font-weight: 900; margin-top: 5px;">{box_counts.get(cat_protect, 0)}</span>
                    </div>
                </div>

                <!-- MIDDLE ROW: MID PEOPLE SCORE -->
                <div style="display: flex; height: 120px;">
                    <div style="flex: 0.3; background-color: #dbe4f0; display: flex; align-items: center; justify-content: center; font-weight: bold; border-right: 2px solid white; border-bottom: 2px solid white;">MID</div>
                    <div style="flex: 1; background-color: {color_map[cat_reduce_some]}; padding: 10px; border-right: 2px solid white; border-bottom: 2px solid white; display: flex; flex-direction: column; justify-content: center;">
                        <span style="font-size: 18px; font-weight: bold;">{cat_reduce_some}</span>
                        <span style="font-size: 24px; font-weight: 900; margin-top: 5px;">{box_counts.get(cat_reduce_some, 0)}</span>
                    </div>
                    <div style="flex: 1; background-color: {color_map[cat_simplify]}; padding: 10px; border-right: 2px solid white; border-bottom: 2px solid white; display: flex; flex-direction: column; justify-content: center;">
                        <span style="font-size: 18px; font-weight: bold;">{cat_simplify}</span>
                        <span style="font-size: 24px; font-weight: 900; margin-top: 5px;">{box_counts.get(cat_simplify, 0)}</span>
                    </div>
                    <div style="flex: 1; background-color: {color_map[cat_share]}; padding: 10px; border-bottom: 2px solid white; display: flex; flex-direction: column; justify-content: center;">
                        <span style="font-size: 18px; font-weight: bold;">{cat_share}</span>
                        <span style="font-size: 24px; font-weight: 900; margin-top: 5px;">{box_counts.get(cat_share, 0)}</span>
                    </div>
                </div>

                <!-- BOTTOM ROW: LOW PEOPLE SCORE -->
                <div style="display: flex; height: 120px;">
                    <div style="flex: 0.3; background-color: #dbe4f0; display: flex; align-items: center; justify-content: center; font-weight: bold; border-right: 2px solid white;">LOW</div>
                    <div style="flex: 1; background-color: {color_map[cat_reduce_most]}; color: white; padding: 10px; border-right: 2px solid white; display: flex; flex-direction: column; justify-content: center;">
                        <span style="font-size: 18px; font-weight: bold;">{cat_reduce_most}</span>
                        <span style="font-size: 24px; font-weight: 900; margin-top: 5px;">{box_counts.get(cat_reduce_most, 0)}</span>
                    </div>
                    <div style="flex: 1; background-color: {color_map[cat_replace]}; padding: 10px; border-right: 2px solid white; display: flex; flex-direction: column; justify-content: center;">
                        <span style="font-size: 18px; font-weight: bold;">{cat_replace}</span>
                        <span style="font-size: 24px; font-weight: 900; margin-top: 5px;">{box_counts.get(cat_replace, 0)}</span>
                    </div>
                    <div style="flex: 1; background-color: {color_map[cat_upskill]}; padding: 10px; display: flex; flex-direction: column; justify-content: center;">
                        <span style="font-size: 18px; font-weight: bold;">{cat_upskill}</span>
                        <span style="font-size: 24px; font-weight: 900; margin-top: 5px;">{box_counts.get(cat_upskill, 0)}</span>
                    </div>
                </div>
            </div>
        </div>
    </div>
    """.replace('\n', '')

    st.markdown(html_matrix, unsafe_allow_html=True)

    # ---------------------------------------------------------
    # DOWNLOAD BUTTON FOR MATRIX SUMMARY (PPTX)
    # ---------------------------------------------------------
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN
        from pptx.enum.text import MSO_ANCHOR

        prs = Presentation()
        # Use a generic layout that is usually completely blank
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Manually add the title box
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9), Inches(0.8))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = "Strategic 9-Box Matrix Summary"
        p.font.bold = True
        p.font.size = Pt(28)
        p.font.color.rgb = RGBColor(30, 58, 95)

        # Define table dimensions and position
        rows, cols = 5, 5
        left, top, width, height = Inches(0.5), Inches(1.3), Inches(9), Inches(4.5)
        table_shape = slide.shapes.add_table(rows, cols, left, top, width, height)
        table = table_shape.table

        # Column widths
        table.columns[0].width = Inches(0.8)
        table.columns[1].width = Inches(0.8)
        table.columns[2].width = Inches(2.46)
        table.columns[3].width = Inches(2.46)
        table.columns[4].width = Inches(2.46)

        # Merge cells for headers
        table.cell(0, 0).merge(table.cell(0, 1))
        table.cell(0, 2).merge(table.cell(0, 4))
        table.cell(1, 0).merge(table.cell(4, 0))


        def hex_to_rgb(hex_str):
            hex_str = hex_str.lstrip('#')
            return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))


        color_lookup = {
            cat_protect: hex_to_rgb('a7c957'),
            cat_keep: hex_to_rgb('c5e1a5'),
            cat_move: hex_to_rgb('a8dadc'),
            cat_share: hex_to_rgb('ffd166'),
            cat_simplify: hex_to_rgb('ffe66d'),
            cat_reduce_some: hex_to_rgb('f4a261'),
            cat_upskill: hex_to_rgb('f4a261'),
            cat_replace: hex_to_rgb('e76f51'),
            cat_reduce_most: hex_to_rgb('d62828')
        }

        # Insert Data directly into origin cells
        # Top-Left Origin (0,0)
        c00 = table.cell(0, 0)
        c00.text = f"TOTAL: {total_persons} {total_person_text.upper()}"
        c00.fill.solid()
        c00.fill.fore_color.rgb = hex_to_rgb('0f172a')

        # Top-Right Origin (0,2)
        c02 = table.cell(0, 2)
        c02.text = f"FINANCIAL OUTPUT ({short_x})"
        c02.fill.solid()
        c02.fill.fore_color.rgb = hex_to_rgb('1e3a5f')

        # Left-Middle Origin (1,0)
        c10 = table.cell(1, 0)
        c10.text = f"PEOPLE SCORE\n({y_metric})"
        c10.fill.solid()
        c10.fill.fore_color.rgb = hex_to_rgb('1e3a5f')

        # Headers Row (1,1 to 1,4)
        h_r1 = ["", "LOW", "MID", "HIGH"]
        for col_idx in range(1, 5):
            c = table.cell(1, col_idx)
            c.text = h_r1[col_idx - 1]
            c.fill.solid()
            c.fill.fore_color.rgb = hex_to_rgb('e6e6e6')

        # Data Rows
        row_labels = ["HIGH", "MID", "LOW"]
        matrix_cats = [
            [cat_move, cat_keep, cat_protect],
            [cat_reduce_some, cat_simplify, cat_share],
            [cat_reduce_most, cat_replace, cat_upskill]
        ]

        origin_cells = [(0, 0), (0, 2), (1, 0)]

        for r_idx, r in enumerate(range(2, 5)):
            c_head = table.cell(r, 1)
            c_head.text = row_labels[r_idx]
            c_head.fill.solid()
            c_head.fill.fore_color.rgb = hex_to_rgb('dbe4f0')
            origin_cells.append((r, 1))

            for c_idx, c in enumerate(range(2, 5)):
                cat_name = matrix_cats[r_idx][c_idx]
                count = box_counts.get(cat_name, 0)
                cell = table.cell(r, c)
                cell.text = f"{cat_name}\n{count}"
                cell.fill.solid()
                cell.fill.fore_color.rgb = color_lookup.get(cat_name, hex_to_rgb('ffffff'))
                origin_cells.append((r, c))

        # Header Row Empty Space
        origin_cells.extend([(1, 1), (1, 2), (1, 3), (1, 4)])

        # Apply Global Formatting
        for r, c in origin_cells:
            cell = table.cell(r, c)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                # Size logic
                if r == 0 or (r == 1 and c == 0):
                    paragraph.font.size = Pt(12)
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = RGBColor(255, 255, 255)
                elif r == 1 or c == 1:
                    paragraph.font.size = Pt(11)
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = RGBColor(0, 0, 0)
                else:
                    paragraph.font.size = Pt(11)
                    paragraph.font.bold = True
                    paragraph.font.color.rgb = RGBColor(31, 41, 55)

        ppt_stream = io.BytesIO()
        prs.save(ppt_stream)
        ppt_stream.seek(0)

        st.download_button(
            label="📥 Download 3x3 Matrix Summary (PowerPoint)",
            data=ppt_stream,
            file_name="Marketing_9Box_Summary.pptx",
            mime="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            type="primary",
            key="btn_dl_summary_ppt"
        )
    except ImportError:
        st.warning(
            "⚠️ To enable PowerPoint download, please install the required library via terminal: `pip install python-pptx`")

    # ==========================================
    # DATA TABLE RESULTS & EXCEL EXPORT
    # ==========================================
    st.markdown("---")
    st.markdown("### 📋 Detailed Personnel Data")

    export_cols = ['Name', y_metric, 'Contribution Margin', 'Operating Profit']
    if x_metric == 'Adjusted Contribution Margin':
        export_cols.append('Adjusted Contribution Margin')
    export_cols.extend(['Sales', '9-Box Category'])

    df_display = df_res[export_cols].copy()
    df_display.rename(columns={'Sales': 'Sales (Bubble Size)'}, inplace=True)
    df_display = df_display.sort_values('9-Box Category')

    format_dict = {
        y_metric: '{:.2f}',
        'Contribution Margin': 'Rp {:,.0f}',
        'Operating Profit': 'Rp {:,.0f}'
    }
    if x_metric == 'Adjusted Contribution Margin':
        format_dict['Adjusted Contribution Margin'] = 'Rp {:,.0f}'
    elif x_metric not in format_dict:
        format_dict[x_metric] = 'Rp {:,.0f}'
    format_dict['Sales (Bubble Size)'] = 'Rp {:,.0f}'

    buffer_data = io.BytesIO()
    with pd.ExcelWriter(buffer_data, engine='openpyxl') as writer:
        df_display.to_excel(writer, index=False, sheet_name='Data_Result')

    st.download_button(
        label="📥 Download Detailed Data (Excel)",
        data=buffer_data.getvalue(),
        file_name="Marketing_Data_Results.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="secondary",
        key="btn_dl_data"
    )

    st.dataframe(df_display.style.format(format_dict), use_container_width=True, hide_index=True)

except Exception as e:
    st.error(f"⚠️ An error occurred while processing the Excel file. Error Log: {e}")
